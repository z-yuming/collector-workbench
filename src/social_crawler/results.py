"""Import old histories on change and page/filter the unified SQLite index."""
from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
import threading
import time
from itertools import islice
from contextlib import closing
from datetime import date, timedelta, datetime
from pathlib import Path
from urllib.parse import quote

from social_crawler import workbench
from social_crawler.settings import Settings
from social_crawler.task_store import TaskStore

_lock = threading.RLock()
_media_scan = {}


def status_of(row):
    if row.get("status") in {"queued", "running", "cancelling", "cancelled", "success", "partial", "failed", "interrupted"}:
        return row["status"]
    return "partial" if row.get("quality", {}).get("status") == "warning" else "success" if row.get("ok") else "failed"


def fingerprint(path):
    if not path.exists():
        return "missing"
    stat = path.stat()
    return f"{stat.st_mtime_ns}:{stat.st_size}"


def unchanged(store, name, stamp):
    with closing(store.connect()) as db:
        row = db.execute("SELECT fingerprint FROM imports WHERE name=?", (name,)).fetchone()
        return row and row[0] == stamp


def replace_import(store, name, stamp, rows):
    with closing(store.connect()) as db:
        db.execute("CREATE TABLE IF NOT EXISTS import_items(name TEXT, id TEXT, PRIMARY KEY(name,id))")
        db.execute("BEGIN IMMEDIATE")
        db.execute("DELETE FROM records WHERE id IN (SELECT id FROM import_items WHERE name=?)", (name,))
        db.execute("DELETE FROM import_items WHERE name=?", (name,))
        for row in rows:
            row["status"] = status_of(row)
            TaskStore.put_record(db, row)
            db.execute("INSERT OR IGNORE INTO import_items VALUES(?,?)", (name, row["id"]))
        db.execute("INSERT OR REPLACE INTO imports VALUES(?,?)", (name, stamp))
        db.commit()


def read_jsonl(path, source):
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines() if path.exists() else []:
        try:
            row = json.loads(line)
            if isinstance(row, dict) and row.get("id"):
                rows.append({**row, "source": source})
        except json.JSONDecodeError:
            continue
    return rows


def table_data(data):
    if not isinstance(data, list):
        return []
    if data and all(isinstance(row, dict) and "items" in row and "key" in row for row in data):
        columns = {str(row.get("name") or row["key"]): row.get("items", []) for row in data}
        return [{name: items[index] if index < len(items) else "" for name, items in columns.items()} for index in range(min(1000, max((len(items) for items in columns.values()), default=0)))]
    return [row for row in data if isinstance(row, dict)]


def scraper_rows(database):
    records = []
    histories = {}
    with closing(sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)) as db:
        db.row_factory = sqlite3.Row
        for table in ("scrape_history", "crawler_jobs"):
            for raw in db.execute(f"SELECT * FROM {table} ORDER BY id DESC"):
                row = dict(raw)
                status = {"success": "success", "completed": "success", "done": "success", "running": "running", "pending": "queued", "queued": "queued", "stopped": "cancelled", "cancelled": "cancelled", "partial": "partial"}.get(row.get("status"), "failed")
                try:
                    data = table_data(json.loads(row.get("data_json") or "[]"))
                except (ValueError, TypeError):
                    data = []
                identity = tuple(row.get(k) or "" for k in ("platform", "url", "keyword", "mode", "shop_name", "shop_id"))
                if table == "scrape_history" and data:
                    histories.setdefault(identity, []).append((row.get("created_at", ""), data))
                if table == "crawler_jobs" and not data:
                    try:
                        job_time = datetime.fromisoformat(row.get("created_at") or "")
                        related = [(abs((datetime.fromisoformat(created) - job_time).total_seconds()), items) for created, items in histories.get(identity, [])]
                        related.sort(key=lambda candidate: candidate[0])
                        if related and related[0][0] <= 60:
                            data = related[0][1]
                    except ValueError:
                        pass
                records.append({"id": f"scraper-{table}-{row['id']}", "source": "scraper", "status": status, "createdAt": row.get("created_at", ""), "platform": row.get("platform", ""), "mode": row.get("mode", ""), "taskType": row.get("task_type") or "products", "query": row.get("keyword") or row.get("url") or row.get("shop_name") or "店铺任务", "ok": status in {"success", "partial"}, "quality": {"status": "ok" if status == "success" else "warning" if status in {"partial", "queued", "running"} else "failed", "label": {"success": "完成", "partial": "需确认", "queued": "排队中", "running": "处理中", "cancelled": "已取消"}.get(status, "失败")}, "counts": {"contents": row.get("item_count", 0), "comments": 0}, "message": row.get("error") or "", "downloadLinks": {}, "preview": {}, "data": data, "originalFile": row.get("output_file") or "", "originalSource": table})
    return records


def sync_history(settings):
    store = TaskStore.for_settings(settings)
    warnings = []
    with _lock:
        for name, filename, source in (("collect-jsonl", "task_history.jsonl", "collect"), ("video-jsonl", "workbench-video-history.jsonl", "video")):
            path = settings.log_dir / filename
            stamp = fingerprint(path)
            if not unchanged(store, name, stamp):
                try:
                    replace_import(store, name, stamp, read_jsonl(path, source))
                except (OSError, ValueError):
                    warnings.append(f"{source}历史暂时无法读取，已保留上次索引。")
        database = workbench.SCRAPER_ROOT / "data.db"
        stamp = "v2:" + fingerprint(database) + ":" + fingerprint(Path(str(database) + "-wal"))
        if not unchanged(store, "scraper-db", stamp):
            try:
                replace_import(store, "scraper-db", stamp, scraper_rows(database) if database.exists() else [])
            except sqlite3.Error:
                warnings.append("店铺与竞品历史暂时无法读取，已保留上次索引。")
        key = str(store.path.resolve())
        if time.monotonic() - _media_scan.get(key, -60) >= 60:
            try:
                root = (workbench.VIDEO_ROOT / "downloads").resolve()
                with closing(store.connect()) as db:
                    known = {link for row in db.execute("SELECT document FROM records WHERE source='video' AND id NOT LIKE 'saved-video-%'") for link in json.loads(row[0]).get("downloadLinks", {}).values()}
                rows = []
                for path in root.rglob("*") if root.exists() else []:
                    if path.is_file() and path.suffix.lower() in {".mp4", ".webm", ".mkv", ".mov", ".mp3", ".m4a", ".jpg", ".jpeg", ".png", ".webp"} and path.resolve().is_relative_to(root):
                        relative = path.resolve().relative_to(root).as_posix()
                        link = "/api/workbench/video-file/" + quote(relative, safe="/")
                        if path.stat().st_size and link not in known:
                            rows.append({"id": "saved-video-" + relative, "source": "video", "status": "success", "createdAt": datetime.fromtimestamp(path.stat().st_mtime).isoformat(), "query": path.name, "platform": "douyin", "taskType": "saved_media", "ok": True, "quality": {"status": "ok", "label": "已保存"}, "downloadLinks": {path.name: link}, "counts": {"contents": 1}, "preview": {}})
                replace_import(store, "saved-media", str(time.time()), rows)
                _media_scan[key] = time.monotonic()
            except OSError:
                warnings.append("部分视频文件暂时无法读取。")
    return store, warnings


def list_results(settings, page=0, limit=20, *, source="all", search="", status="all", platform="", start="", end=""):
    if source not in {"all", "collect", "scraper", "video"}:
        raise ValueError("任务来源无效。")
    if status not in {"all", "queued", "running", "cancelling", "success", "partial", "failed", "cancelled", "interrupted"}:
        raise ValueError("任务状态无效。")
    if page < 0 or not 1 <= limit <= 100:
        raise ValueError("任务记录页码或每页数量无效。")
    clauses = []; params = []
    for column, value in (("source", source), ("status", status), ("platform", platform)):
        if value and value != "all":
            clauses.append(f"{column}=?"); params.append(value)
    if search:
        clauses.append("(query LIKE ? ESCAPE '\\' OR platform LIKE ? ESCAPE '\\')")
        term = "%" + search.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
        params.extend([term, term])
    try:
        if start:
            clauses.append("created>=?"); params.append(date.fromisoformat(start).isoformat())
        if end:
            clauses.append("created<?"); params.append((date.fromisoformat(end) + timedelta(days=1)).isoformat())
        if start and end and start > end:
            raise ValueError()
    except ValueError:
        raise ValueError("日期范围无效，请使用有效的开始和结束日期。")
    store, warnings = sync_history(settings)
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with closing(store.connect()) as db:
        total = db.execute("SELECT COUNT(*) FROM records" + where, params).fetchone()[0]
        rows = db.execute("SELECT document FROM records" + where + " ORDER BY created DESC,id DESC LIMIT ? OFFSET ?", [*params, limit, page * limit])
        items = [{k: v for k, v in json.loads(row[0]).items() if k not in {"data", "settings", "exports", "validationOutputs", "originalFile"}} for row in rows]
        counts = dict.fromkeys(("collect", "scraper", "video"), 0)
        counts.update({row[0]: row[1] for row in db.execute("SELECT source,COUNT(*) FROM records GROUP BY source")})
    return {"items": items, "total": total, "page": page, "limit": limit, "sourceCounts": counts, "warnings": warnings}


def result_detail(settings, record_id):
    store, _ = sync_history(settings)
    with closing(store.connect()) as db:
        row = db.execute("SELECT document FROM records WHERE id=?", (record_id,)).fetchone()
    if not row:
        raise KeyError("任务记录不存在，请刷新列表。")
    detail = json.loads(row[0])
    if detail.get("source") != "scraper":
        return detail
    data = detail.pop("data", [])
    original = Path(detail.pop("originalFile", ""))
    allowed_roots = [settings.output_dir.resolve(), (workbench.SCRAPER_ROOT / "crawler_output").resolve(), (workbench.SCRAPER_ROOT / "media_crawler/data").resolve()]
    if original.is_file() and original.suffix.lower() in {".xlsx", ".csv", ".json", ".jsonl", ".mp4", ".webm", ".jpg", ".png"} and any(original.resolve().is_relative_to(root) for root in allowed_roots):
        if original.suffix.lower() == ".xlsx" and not data:
            from openpyxl import load_workbook
            book = load_workbook(original, read_only=True, data_only=True)
            try:
                sheet = book["contents"] if "contents" in book.sheetnames else book.active
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(x or "字段") for x in next(iterator, [])]
                data = [dict(zip(headers, row)) for row in islice(iterator, 1000)]
            finally:
                book.close()
        detail["originalDownload"] = True
        detail["downloadLinks"]["原始文件"] = "/api/workbench/result-file/" + quote(record_id, safe="")
    if data:
        columns = list(dict.fromkeys(key for item in data for key in item))[:30]
        detail["preview"] = {"table": {"columns": columns, "rows": data[:50]}, "total": len(data)}
        stem = "history_" + hashlib.sha256((record_id + json.dumps(data, ensure_ascii=False, default=str)).encode()).hexdigest()[:20]
        settings.output_dir.mkdir(parents=True, exist_ok=True)
        for extension in ("json", "csv", "xlsx"):
            path = settings.output_dir / (stem + "." + extension)
            if not path.exists():
                if extension == "json":
                    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
                elif extension == "csv":
                    with path.open("w", newline="", encoding="utf-8-sig") as output:
                        writer = csv.writer(output); writer.writerow(columns)
                        writer.writerows([("'" + str(value)) if str(value).startswith(("=", "+", "-", "@")) else value for value in (item.get(k, "") for k in columns)] for item in data)
                else:
                    from openpyxl import Workbook
                    book = Workbook(); sheet = book.active; sheet.title = "结果"
                    sheet.append(columns)
                    for item in data:
                        sheet.append([str(item.get(k, "")) if item.get(k) is not None else "" for k in columns])
                        for cell in sheet[sheet.max_row]:
                            cell.data_type = "s"
                    book.save(path)
            detail["downloadLinks"]["excel" if extension == "xlsx" else extension] = "/api/file/" + path.name
    elif not detail["downloadLinks"]:
        detail["message"] = detail.get("message") or "旧记录没有保存可预览数据或可用文件，原任务记录仍保留。"
    return detail


def original_result_file(settings, record_id):
    store, _ = sync_history(settings)
    with closing(store.connect()) as db:
        row = db.execute("SELECT document FROM records WHERE id=? AND source='scraper'", (record_id,)).fetchone()
    if row:
        path = Path(json.loads(row[0]).get("originalFile") or "")
        roots = [settings.output_dir.resolve(), (workbench.SCRAPER_ROOT / "crawler_output").resolve(), (workbench.SCRAPER_ROOT / "media_crawler/data").resolve()]
        if path.is_file() and path.suffix.lower() in {".xlsx", ".csv", ".json", ".jsonl", ".mp4", ".webm", ".jpg", ".png"} and any(path.resolve().is_relative_to(root) for root in roots):
            return path
    raise ValueError("原文件不存在或保存位置不在结果目录内。")
